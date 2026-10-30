import sys,json,csv,re,collections,hashlib,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'.deps'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from config import ROOT,DATE,TOPICS,SEEDS,ALIASES,EUROPE

OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
inventory=json.loads((ROOT/'inventory.json').read_text());rows=inventory['universities']
scans=[json.loads(l) for l in (ROOT/'portal_scan.jsonl').read_text().split('\n') if l]
scanmap={r['record_id']:r for r in scans}
def canonical(d):
 d=d.lower().removeprefix('www.')
 if d=='daihoc.fpt.edu.vn':return 'fpt.edu.vn'
 return ALIASES.get(d,d)
seedmap={canonical(d):(i,name,country) for i,(name,country,d) in enumerate(SEEDS)}
schools=[]
for part in ['vn_core_research','vn_research','us_research','eu_research']:
 for s in json.loads((ROOT/'parts'/f'{part}.json').read_text())['schools']:
  s=dict(s);key=canonical(s['domain']);order,short,country=seedmap[key]
  s.update(seed_order=order,short_name=short,country_code=country,canonical_domain=key)
  s.setdefault('website','https://'+s['domain']+'/')
  s['region']='Việt Nam' if country=='VN' else 'Hoa Kỳ' if country=='US' else 'Châu Âu mở rộng'
  s['research_scope_note']='Các công trình được chọn có nguồn; không phải toàn bộ sản lượng hay mọi hướng nghiên cứu của trường.'
  schools.append(s)
schools.sort(key=lambda s:s['seed_order'])
works=[]
for s in schools:
 for item in s['items']:
  i=dict(item)
  i['school']=s['name'];i['short_name']=s['short_name'];i['school_domain']=s['domain'];i['country_code']=s['country_code'];i['region']=s['region']
  i['record_id']='R'+str(len(works)+1).zfill(3)
  i['topic_codes']=list(dict.fromkeys(i['topic_codes']))
  i.setdefault('primary_topic_code',i['topic_codes'][0]);i['secondary_topic_codes']=[c for c in i['topic_codes'] if c!=i['primary_topic_code']]
  i['primary_topic_vi']=TOPICS[i['primary_topic_code']][0]
  i['topic_labels_vi']=[TOPICS[c][0] for c in i['topic_codes']]
  i.setdefault('source_urls',[i['url']]);i['source_urls']=list(dict.fromkeys([i['url']]+i['source_urls']))
  i.setdefault('full_text_read',False);i['checked_on']=DATE
  note=i.get('verification_note','').lower()
  i['evidence_level']='bibliography_only' if any(x in note for x in ['chưa đọc abstract','thông tin thư mục và nhãn']) else 'source_content_reviewed'
  i['review_status']='source_reviewed_with_limitations'
  works.append(i)
schoolmap={s['canonical_domain']:s for s in schools}
for r in rows:
 matches=[schoolmap[canonical(d)] for d in r['domains'] if canonical(d) in schoolmap]
 if matches:
  s=matches[0];r['research_scan_status']='selected_works_reviewed';r['reviewed_research_domain']=s['domain'];r['reviewed_item_count']=len(s['items'])
  r['technology_evidence']='Có công trình công nghệ-kỹ thuật trong tập mẫu đã đọc nguồn; không khẳng định trường chỉ đào tạo công nghệ.'
 else:r['research_scan_status']=scanmap.get(r['record_id'],{}).get('scan_status','not_surveyed')
 r['portal_survey_status']=scanmap.get(r['record_id'],{}).get('scan_status','not_surveyed')

counts=collections.Counter(w['primary_topic_code'] for w in works)
multicounts=collections.Counter(c for w in works for c in w['topic_codes'])
stats={'date':DATE,'directory_records':len(rows),'countries_and_territories':len(set(r['country_code'] for r in rows)),
 'directory_regions':dict(collections.Counter(r['region'] for r in rows)),
 'technology_candidates_worldwide':sum(r['technology_priority'] for r in rows),
 'technology_candidates_priority_surveyed':len(scans),'survey_status_counts':dict(collections.Counter(x['scan_status'] for x in scans)),
 'page_attempts':sum(len(x.get('pages',[])) for x in scans),
 'discovered_portal_links':sum(len(x.get('research_portal_leads',[])) for x in scans),
 'unreviewed_item_link_leads':sum(len(x.get('research_item_leads',[])) for x in scans),
 'schools_with_selected_works':len(schools),'selected_work_records':len(works),
 'work_regions':dict(collections.Counter(w['region'] for w in works)),
 'work_types':dict(collections.Counter(w['type'] for w in works)),
 'primary_topic_counts':dict(counts),'multi_label_topic_counts':dict(multicounts),
 'work_evidence_levels':dict(collections.Counter(w['evidence_level'] for w in works)),
 'works_with_doi':sum(bool(w.get('doi')) for w in works)}
bio=json.loads((ROOT/'bibliography_classified.json').read_text());stats['bibliography_records_automatic_classification']=len(bio['records'])

limitations=[
 'Đây là lượt rà soát ban đầu, chưa hoàn thành việc quét sâu mọi trường trên thế giới hoặc toàn bộ 589 trường trong bảng ban đầu. Tệp nguồn 589/5.021 trường chưa được cung cấp; không thể xác nhận hoặc tái lập các tổng số đó.',
 'Danh bạ gồm bản ghi trường đại học, học viện, college, hệ thống, trường thành viên và một số phân hiệu/lịch sử. Tổng bản ghi không phải số trường đại học pháp lý độc lập; việc gộp tên miền có thể gộp campus, và khác tên/khác domain có thể còn trùng.',
 'Ứng viên công nghệ được chọn theo ví dụ người dùng và từ khóa tên đa ngôn ngữ. Đây là bộ lọc phục vụ ưu tiên; có thể bỏ sót trường đa ngành mạnh về kỹ thuật và giữ nhầm cơ sở không phù hợp. is_technology để null khi chưa thẩm định tư cách chuyên ngành.',
 'Châu Âu dùng phạm vi mở rộng từ EHEA/DEQAR, có UK, Thụy Sĩ, Na Uy, các nước xuyên lục địa, Caucasus và Kazakhstan. Cột ISO2 giúp giới hạn lại theo nhu cầu; Kosovo dùng mã XK do nguồn gán.',
 'VN-NARIC có lỗi ghép tên–domain, bản ghi cũ và trùng; 48 cặp sai rõ ràng đã bị loại. Trích xuất VN-NARIC gặp chứng thư TLS không khớp hostname, được ghi trong nhật ký nguồn. Không coi nguồn danh bạ là chứng minh website còn chính thức hoặc trường được kiểm định hiện hành.',
 'Nguồn Mỹ chủ yếu là danh bạ cộng đồng Hipo. NCES HD2025 được tìm thấy nhưng máy chủ tải dữ liệu không truy cập được. Bản ROR tải không hoàn chỉnh nên không dùng để tạo các bản ghi.',
 'Khảo sát tự động giới hạn khoảng 2–4 yêu cầu trang mỗi bản ghi, ưu tiên trang nghiên cứu/công bố cùng tên miền. Chỉ phân tích HTML công khai, không vượt đăng nhập hoặc đọc toàn bộ PDF; lỗi TLS, chống bot, JavaScript và timeout làm giảm độ phủ. HTTP200 cũng không chứng minh danh tính website.',
 'Liên kết tự động là đầu mối cần kiểm tra, không phải bài báo đã xác nhận. Nhãn tự động của 366 bản ghi thư mục dựa vào tiêu đề, có thể sai; không dùng để suy ra thế mạnh hay xếp hạng trường.',
 '87 mục chi tiết là mẫu công trình có nguồn thuộc 41 trường. Một phần chỉ xác nhận thông tin thư mục, số khác đọc abstract hoặc bản tin cơ quan; chưa thẩm định toàn văn, phương pháp nghiên cứu hoặc khả năng tái lập. Một số trường hợp quy thuộc trường chỉ qua kho cơ quan/email tác giả và đã ghi chú.',
 'Bài đã chấp nhận, bài công bố, báo cáo kỹ thuật, luận án, đồ án và dự án được giữ loại riêng. Không coi dự án đang triển khai là kết quả đã công bố. Khi năm hội nghị và năm phát hành khác nhau, giữ ghi chú; năm chưa rõ để trống.',
 'Một công trình có thể thuộc nhiều chủ đề; tổng số nhãn sẽ vượt số bản ghi. Không cộng trực tiếp 87 mục chi tiết với 366 bản ghi thư mục vì có giao nhau. Không lấy số mẫu để kết luận tỷ trọng công bố của toàn trường.',
]
meta={**inventory['metadata'],'scan_stats':stats,'limitations_vi':limitations,'topic_taxonomy':{k:v[0] for k,v in TOPICS.items()},
 'sources_frozen':'Dữ liệu được lấy ngày2026-10-30; DEQAR daily CSV last-modified2026-10-30; Hipo master snapshot lưu cục bộ kèmSHA256.',
 'raw_source_sha256':{'hipo.json':hashlib.sha256((ROOT/'raw/hipo.json').read_bytes()).hexdigest()}}

def dump(name,obj):
 (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
dump('universities_tech_focus.json',{'metadata':meta,'technology_universities_with_reviewed_works':schools,
 'technology_candidates':[r for r in rows if r['technology_priority']],'all_universities':rows})
dump('research_classified.json',{'metadata':{'date':DATE,'statistics':stats,'limitations_vi':limitations,'topic_taxonomy':{k:v[0] for k,v in TOPICS.items()}},'schools':schools,'research_items':works})
dump('bibliography_classified.json',bio)
dump('portal_scan.json',{'metadata':{'date':DATE,'method':'Khảo sát HTML có giới hạn; các liên kết tìm thấy chưa được xác minh là công trình của trường.','statistics':stats},'records':scans})
dump('source_notes.json',{'limitations_vi':limitations,'regional_notes':{n:json.loads((ROOT/'parts'/n).read_text()) for n in ['vietnam_notes.json','europe_notes.json','usa_notes.json']}})
dump('statistics.json',stats)

types={
 'journal_article':'Bài báo tạp chí','conference_paper':'Bài hội nghị','accepted_conference_paper':'Bài hội nghị đã chấp nhận; chưa xác nhận xuất bản',
 'book_chapter':'Chương sách/kỷ yếu','technical_report':'Báo cáo kỹ thuật','doctoral_dissertation':'Luận án tiến sĩ','doctoral_thesis':'Luận án tiến sĩ',
 'student_graduation_project':'Đồ án tốt nghiệp sinh viên','institutional_research_project':'Đề tài/dự án của cơ sở',
 'funded_research_project':'Đề tài/dự án được tài trợ','research_project':'Dự án nghiên cứu','review_article':'Bài tổng quan'}
status_vi={'research_links_found':'Tìm được liên kết nghiên cứu','homepage_only_no_research_link_found':'Đọc trang chính; chưa tìm thấy liên kết nghiên cứu',
 'fetch_failed':'Không lấy được nội dung','not_surveyed':'Chưa khảo sát','selected_works_reviewed':'Có công trình đã đọc nguồn'}
def tidy(text):
 text=str(text or '')
 text=re.sub(r'(?<=[^\W\d_])(?=20[12]\d\b)',' ',text)
 text=re.sub(r'(?<=[.,;])(?=[A-Za-zÀ-ỹ])',' ',text)
 for a,b in [('mạng5G','mạng 5G'),('thông1','thông 1'),('tác giảDUT','tác giả DUT'),('tác giảChungnam','tác giả Chungnam'),('thuộcHCMUTE','thuộc HCMUTE'),('khoVKU','kho VKU'),('báoVKU','báo VKU'),('trên websiteKhoa','trên website Khoa')]:text=text.replace(a,b)
 return text
def flat(v):
 if v is None:return ''
 if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False) if isinstance(v,dict) else '; '.join(map(str,v))
 return v
def savecsv(name,headers,data):
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f);w.writerow(headers)
  for row in data:w.writerow([flat(v) for v in row])

work_headers=['ID','Trường','ISO2','Website trường','Tên công trình','Năm','Loại','Chủ đề chính','Chủ đề phụ','Từ khóa chi tiết','Tóm tắt tiếng Việt','DOI','URL nguồn','Nguồn bổ sung','Mức chứng cứ','Ghi chú kiểm chứng','Đọc toàn văn']
work_data=[[w['record_id'],w['school'],w['country_code'],'https://'+w['school_domain'],w['title'],w['year'],types.get(w['type'],w['type']),w['primary_topic_vi'],[TOPICS[c][0] for c in w['secondary_topic_codes']],w.get('topics',[]),tidy(w['summary_vi']),w.get('doi'),w['url'],w['source_urls'],w['evidence_level'],tidy(w.get('verification_note')),'Chưa'] for w in works]
dir_headers=['ID','Trường','Quốc gia','ISO2','Khu vực','Website theo nguồn','Tên miền','Loại cơ sở','Nguồn','URL nguồn','Ưu tiên công nghệ','Căn cứ chọn','Trạng thái nghiên cứu','Trạng thái website','Ghi chú']
dir_data=[[r['record_id'],r['name'],r['country'],r['country_code'],r['region'],r['website'],r['domains'],r.get('institution_type'),r['source_names'],r['source_urls'],r['technology_priority'],r['technology_selection_basis'],status_vi.get(r['research_scan_status'],r['research_scan_status']),r['website_verification_status'],r.get('verification_note')] for r in rows]
savecsv('research_classified.csv',work_headers,work_data)
savecsv('university_websites.csv',dir_headers,dir_data)

wb=Workbook();wb.remove(wb.active)
def sheet(name,headers,data,widths=None):
 ws=wb.create_sheet(name);ws.append(headers)
 for row in data:
  clean=[]
  for val in row:
   v=flat(val)
   if isinstance(v,str):v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:32700]
   clean.append(v)
  ws.append(clean)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183B56');c.alignment=Alignment(wrap_text=True,vertical='top')
 ws.row_dimensions[1].height=32
 for col in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(col)].width=(widths or {}).get(col,25)
 for line in ws.iter_rows(min_row=2):
  for c in line:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(c.value,str) and c.value.startswith(('https://','http://')) and '; ' not in c.value and len(c.value)<1500:
    c.hyperlink=c.value;c.style='Hyperlink';c.alignment=Alignment(vertical='top',wrap_text=True)
 return ws

guide=[['Phạm vi','Ưu tiên trường công nghệ–kỹ thuật ở Việt Nam, Hoa Kỳ và Châu Âu; danh bạ nền có thêm các nước khác.'],
 ['Kết quả',f'{len(rows):,} bản ghi danh bạ; {len(scans):,} ứng viên khảo sát cổng nghiên cứu; {len(schools)} trường có {len(works)} mục đọc nguồn; {len(bio["records"])} bản ghi thư mục phân loại tự động.'],
 ['Cách dùng','Cong_trinh: đọc mẫu đã kiểm tra; Chu_de: lọc ngành; Truong_da_doc: trường và cổng nghiên cứu; Scan_cong_nghe: tiến độ khảo sát; Thu_muc_tu_dong và Link_can_kiem_tra: cần kiểm tra thêm; Danh_ba_the_gioi: danh mục website.']]+[[f'Giới hạn {i+1}',t] for i,t in enumerate(limitations)]
sheet('Huong_dan',['Mục','Nội dung'],guide,{1:24,2:120})
coverage=[]
for reg in ['Việt Nam','Hoa Kỳ','Châu Âu mở rộng','Khu vực khác']:
 rr=[r for r in rows if r['region']==reg];ss=[s for s in schools if s['region']==reg]
 coverage.append([reg,len(rr),sum(r['technology_priority'] for r in rr),sum(r['record_id'] in scanmap for r in rr),len(ss),sum(len(s['items']) for s in ss)])
sheet('Bao_phu',['Khu vực','Bản ghi danh bạ','Ứng viên công nghệ','Đã thử khảo sát cổng','Trường có mục chi tiết','Mục chi tiết'],coverage,{1:26})
sheet('Chu_de',['Mã','Nhóm chủ đề','Số mục: nhãn chính','Số mục: mọi nhãn'],[[k,v[0],counts[k],multicounts[k]] for k,v in TOPICS.items()],{1:30,2:55})
sheet('Cong_trinh',work_headers,work_data,{2:50,5:75,7:35,8:48,9:55,10:45,11:95,13:70,14:80,16:95})
sheet('Truong_da_doc',['Trường','Tên ngắn','ISO2','Website','Cổng nghiên cứu','Số mục','Chủ đề trong mẫu','Phạm vi'],[[s['name'],s['short_name'],s['country_code'],s['website'],[p['url'] for p in s['research_portals']],len(s['items']),sorted({TOPICS[c][0] for i in s['items'] for c in i['topic_codes']}),s['research_scope_note']] for s in schools],{1:60,4:40,5:85,7:65,8:65})
sheet('Scan_cong_nghe',['Trường','ISO2','Website','Căn cứ chọn','Trạng thái','Số lần lấy trang','Cổng nghiên cứu tìm được','Số đầu mối bài/đề tài chưa kiểm tra'],[[s['name'],s['country_code'],s['website'],s['technology_selection_basis'],status_vi.get(s['scan_status'],s['scan_status']),len(s['pages']),[p['url'] for p in s['research_portal_leads']],len(s['research_item_leads'])] for s in scans],{1:60,3:45,5:55,7:85})
leads=[[s['name'],s['country_code'],i['title'],i['url'],i['found_on'],[TOPICS[c][0] for c in i['topic_codes_auto']],'Liên kết tự động; chưa xác nhận loại công trình hoặc tác giả thuộc trường'] for s in scans for i in s['research_item_leads']]
sheet('Link_can_kiem_tra',['Trường','ISO2','Văn bản liên kết','URL','Tìm thấy ở','Nhãn tự động','Kiểm chứng'],leads,{1:50,3:80,4:75,5:75,6:60,7:70})
sheet('Thu_muc_tu_dong',['Trường','Tên công trình','Năm','Loại bản ghi','Nhãn theo tiêu đề','URL nguồn','URL công bố','DOI','Tác giả theo nguồn','Ghi chú'],[[r['school'],r['title'],r['year'],r['type'],r['topics'],r['source_url'],r['publication_urls'],r.get('doi'),r['authors_as_listed'],r['classification_method']+' '+r['verification_note']] for r in bio['records']],{2:90,5:65,6:70,7:70,9:65,10:100})
sheet('Danh_ba_the_gioi',dir_headers,dir_data,{2:65,6:50,7:40,9:45,10:80,13:55,15:95})
sheet('Nguon',['Nguồn','URL','Vai trò'],[
 ['VN-NARIC','https://en.naric.edu.vn/vietnam-recognised-bodies','Danh bạ cơ sở Việt Nam; đã loại các cặp sai rõ ràng'],
 ['EQAR/DEQAR','https://www.eqar.eu/qa-results/download-data-sets/','Danh bạ cơ sở Châu Âu; dữ liệu tải theo ngày'],
 ['DEQAR CSV','https://backend.deqar.eu/static/daily-csv/deqar-institutions.csv','Tệp dữ liệu nguồn'],
 ['France MESR','https://data.enseignementsup-recherche.gouv.fr/explore/dataset/fr-esr-principaux-etablissements-enseignement-superieur/','Đối chiếu/bổ sung cơ sở Pháp'],
 ['Hipo','https://github.com/Hipo/university-domains-list','Danh bạ cộng đồng toàn cầu; không phải danh bạ kiểm định'],
 ['Nguồn công trình','Xem Cong_trinh và Truong_da_doc','URL nguồn riêng cho từng trường và từng công trình']
 ],{1:25,2:105,3:95})
wb.save(OUT/'bao_cao_truong_cong_nghe.xlsx')

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.header.paragraphs[0].text='KHẢO SÁT NGHIÊN CỨU CÔNG NGHỆ • VIỆT NAM – MỸ – CHÂU ÂU'
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
for style in ['Title','Heading 1','Heading 2','Heading 3']:doc.styles[style].font.color.rgb=RGBColor.from_string('183B56')
def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'0563C1');pr.append(col);run.append(pr)
 t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);p._p.append(h)
def para(text):return doc.add_paragraph(tidy(text))
doc.add_heading('Báo cáo website và phân loại nghiên cứu công nghệ–kỹ thuật',0)
para('Ngày thu thập:30/10/2026. Phiên bản đầu; phạm vi và mức kiểm chứng được nêu dưới đây.')
para(f'Đã tổng hợp {len(rows):,} bản ghi cơ sở giáo dục tại {stats["countries_and_territories"]} quốc gia/vùng lãnh thổ; khảo sát cổng nghiên cứu của {len(scans):,} ứng viên ưu tiên. Phần chi tiết có {len(works)} mục thuộc {len(schools)} trường: Việt Nam 39, Mỹ 20, Châu Âu 28. Số liệu là quy mô dữ liệu đã thu thập, không phải tổng số trường/công bố thực tế.')
doc.add_heading('1. Phạm vi và cách đọc',1)
for t in limitations:para(t)
doc.add_heading('2. Độ phủ dữ liệu',1)
t=doc.add_table(rows=1,cols=5);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Khu vực','Danh bạ','Ứng viên CN','Trường chi tiết','Mục chi tiết']):c.text=v
for reg,total,cands,survey,ns,nw in coverage:
 for c,v in zip(t.add_row().cells,[reg,total,cands,ns,nw]):c.text=str(v)
para(f'Khảo sát tự động: {stats["survey_status_counts"]["research_links_found"]} bản ghi tìm thấy cổng nghiên cứu; {stats["survey_status_counts"]["homepage_only_no_research_link_found"]} chỉ đọc được trang chính; {stats["survey_status_counts"]["fetch_failed"]} không lấy được nội dung. Tổng {stats["page_attempts"]} lần lấy trang, {stats["discovered_portal_links"]} liên kết cổng nghiên cứu và {stats["unreviewed_item_link_leads"]} đầu mối chưa kiểm tra. Các giá trị này không phải số bài báo.')
doc.add_heading('3. Phân loại chủ đề',1)
para('Nhãn chính phản ánh mục tiêu của công trình; nhãn phụ ghi phương pháp hoặc lĩnh vực ứng dụng. Đây là phân loại phục vụ tra cứu của báo cáo, không phải hệ phân ngành chính thức hay xếp hạng trường.')
t=doc.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Nhóm chủ đề','Số mục là nhãn chính','Số mục có nhãn']):c.text=v
for k,v in TOPICS.items():
 for c,x in zip(t.add_row().cells,[v[0],counts[k],multicounts[k]]):c.text=str(x)
doc.add_heading('4. Công trình theo từng trường',1)
for s in schools:
 doc.add_heading(f'{s["seed_order"]+1}. {s["name"]}',2)
 p=doc.add_paragraph();p.add_run('Website: ');link(p,s['website'],s['website']);p.add_run(' | ISO2: '+s['country_code'])
 for portal in s['research_portals']:
  p=doc.add_paragraph();p.add_run('Nguồn tra cứu: ');link(p,tidy(portal['title']),portal['url'])
 sw=[w for w in works if w['school_domain']==s['domain']]
 para('Chủ đề có bằng chứng trong mẫu: '+'; '.join(sorted({TOPICS[c][0] for w in sw for c in w['topic_codes']}))+'.')
 for w in sw:
  doc.add_heading(w['record_id']+' • '+w['title'],3)
  para(f'Năm: {w["year"] if w["year"] else "chưa xác định"} | Loại: {types.get(w["type"],w["type"])}')
  para('Chủ đề chính: '+w['primary_topic_vi']+'. Nhãn bổ sung: '+('; '.join(TOPICS[c][0] for c in w['secondary_topic_codes']) or 'Không')+'.')
  para(w['summary_vi']);para('Mức kiểm chứng: '+w.get('verification_note',''))
  p=doc.add_paragraph();link(p,'Nguồn nội dung',w['url'])
  if w.get('doi'):p.add_run(' | ');link(p,'DOI: '+w['doi'],'https://doi.org/'+w['doi'].removeprefix('https://doi.org/'))
  for url in w['source_urls']:
   if url==w['url']:continue
   p=doc.add_paragraph();link(p,'Nguồn đối chiếu: '+url,url)
doc.add_heading('5. Các tệp dữ liệu kèm theo',1)
for name,description in [
 ('bao_cao_truong_cong_nghe.xlsx','Bộ lọc trường, quốc gia, loại công trình, chủ đề và mức kiểm chứng; đầy đủ danh bạ và nhật ký khảo sát.'),
 ('research_classified.json / .csv','87 mục chi tiết; JSON giữ cấu trúc nhiều nhãn và nhiều nguồn.'),
 ('universities_tech_focus.json','Danh bạ toàn cầu, trường có mục đọc nguồn và ứng viên công nghệ; kèm tiêu chí/giới hạn.'),
 ('bibliography_classified.json','366 bản ghi thư mục phân loại tự động; cần kiểm tra trước khi dùng như dữ liệu chuẩn.'),
 ('portal_scan.json','Nhật ký kết quả và lỗi khi tìm cổng nghiên cứu; đầu mối chưa xác nhận.'),
 ('source_notes.json','Nguồn danh bạ, bộ lọc và các giới hạn riêng từng khu vực.')]:para(name+': '+description)
doc.add_heading('6. Nguồn danh bạ nền',1)
for label,url in [('VN-NARIC','https://en.naric.edu.vn/vietnam-recognised-bodies'),('EQAR/DEQAR','https://www.eqar.eu/qa-results/download-data-sets/'),('France MESR','https://data.enseignementsup-recherche.gouv.fr/explore/dataset/fr-esr-principaux-etablissements-enseignement-superieur/'),('Hipo','https://github.com/Hipo/university-domains-list')]:
 p=doc.add_paragraph();link(p,label,url)
doc.save(OUT/'bao_cao_nghien_cuu_chi_tiet.docx')

with zipfile.ZipFile(ROOT/'bao_cao_truong_cong_nghe.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in OUT.iterdir():z.write(p,p.name)
print(json.dumps(stats,ensure_ascii=False,indent=2))
print('deliverables',[(p.name,p.stat().st_size) for p in OUT.iterdir()])
