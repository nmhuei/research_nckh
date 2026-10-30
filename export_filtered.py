"""Build the technology-only deliverable from explicit relevance decisions."""
import sys,json,re,csv,collections,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'.deps'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from config import ROOT,DATE,TOPICS,SEEDS,ALIASES,classify

OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
AUDITS=['vn_relevance_review','us_relevance_review','eu_relevance_review','hust_bib_relevance','vku_bib_relevance','ptit_bib_relevance']
reviews={n:json.loads((ROOT/'parts'/f'{n}.json').read_text()) for n in AUDITS}
def key(s):return re.sub(r'\s+',' ',s).strip().casefold()
def kept(review):return {key(x if isinstance(x,str) else x['title']) for x in review['keep']}
work_keep=set().union(*(kept(reviews[n]) for n in AUDITS[:3]))
bib_keep=set().union(*(kept(reviews[n]) for n in AUDITS[3:]))
def canonical(d):return 'fpt.edu.vn' if d=='daihoc.fpt.edu.vn' else ALIASES.get(d,d)
seedmap={canonical(d):(i,n,c) for i,(n,c,d) in enumerate(SEEDS)}
schools=[];all_original=[]
for part in ['vn_core_research','vn_research','us_research','eu_research']:
 for school in json.loads((ROOT/'parts'/f'{part}.json').read_text())['schools']:
  all_original.extend(school['items']);s=dict(school)
  s['items']=[dict(w) for w in s['items'] if key(w['title']) in work_keep]
  if not s['items']:continue
  order,short,code=seedmap[canonical(s['domain'])]
  s.update(seed_order=order,short_name=short,country_code=code,region='Việt Nam' if code=='VN' else 'Hoa Kỳ' if code=='US' else 'Châu Âu')
  s.setdefault('website','https://'+s['domain']+'/');s['technology_evidence']='Có công trình công nghệ–kỹ thuật phù hợp trong tập nguồn đã đọc; không đồng nghĩa trường chỉ chuyên ngành công nghệ.'
  schools.append(s)
schools.sort(key=lambda s:s['seed_order'])
taxonomy={k:v[0] for k,v in TOPICS.items() if k not in ['OTHER','BASIC_SCIENCE']}
works=[]
for s in schools:
 for w in s['items']:
  # Precision clocks are instrumentation/quantum technology; avoid placing them
  # in a catch-all basic-science group in the technology-only report.
  w['topic_codes']=list(dict.fromkeys('ELECTRONICS_TELECOM' if c=='BASIC_SCIENCE' else c for c in w['topic_codes'] if c!='OTHER'))
  if w.get('primary_topic_code')=='BASIC_SCIENCE':w['primary_topic_code']='ELECTRONICS_TELECOM'
  w.setdefault('primary_topic_code',w['topic_codes'][0]);w['secondary_topic_codes']=[c for c in w['topic_codes'] if c!=w['primary_topic_code']]
  w.update(record_id='T'+str(len(works)+1).zfill(3),school=s['name'],short_name=s['short_name'],school_domain=s['domain'],country_code=s['country_code'],region=s['region'])
  w['primary_topic_vi']=taxonomy[w['primary_topic_code']];w['topic_labels_vi']=[taxonomy[c] for c in w['topic_codes']]
  w.setdefault('source_urls',[w['url']]);w['source_urls']=list(dict.fromkeys([w['url']]+w['source_urls']))
  w['full_text_read']=False;w['checked_on']=DATE;w['technology_relevance']='included_after_content_relevance_review'
  note=w.get('verification_note','').lower()
  w['evidence_level']='bibliography_only' if any(x in note for x in ['chưa đọc abstract','thông tin thư mục và nhãn']) else 'source_content_reviewed'
  works.append(w)
bio_source=json.loads((ROOT/'bibliography_classified.json').read_text())['records']
bib=[dict(r) for r in bio_source if key(r['title']) in bib_keep]
topic_reviews={key(r['title']):r for name in ['hust_bib_topics','vku_bib_topics','ptit_bib_topics'] for r in json.loads((ROOT/'parts'/f'{name}.json').read_text())['records']}
for r in bib:
 r['technology_relevance']='manually_screened_title_only'
 reviewed_topics=topic_reviews[key(r['title'])]
 r['topic_codes']=reviewed_topics['topic_codes']
 r['primary_topic_code']=reviewed_topics['primary_topic_code']
 r['topics']=[taxonomy[c] for c in r['topic_codes']]
 if not r['topic_codes']:r['topics']=['Có liên quan công nghệ; chưa đủ dữ liệu gán nhóm cụ thể']
 r['classification_method']='Đã rà mức liên quan và nhóm công nghệ theo tiêu đề; một số trường hợp có đối chiếu abstract. Chưa kiểm tra đầy đủ nội dung từng bài hoặc đọc toàn văn.'
 if 'Ứng dụng Lý thuyết Dấu ấn' in r['title']:
  r['verification_note']+=' Lưu ý: danh mục khoa ghi năm 2025; metadata tạp chí VJOL ghi năm 2026. Giữ năm theo danh mục và chưa kết luận năm xuất bản cuối cùng.'
  r['additional_evidence_url']='https://www.vjol.info.vn/3126-3445/article/view/258017/'

counts=collections.Counter(w['primary_topic_code'] for w in works)
allcounts=collections.Counter(c for w in works for c in w['topic_codes'])
coverage=[{'region':region,'schools':sum(s['region']==region for s in schools),'items':sum(w['region']==region for w in works)} for region in ['Việt Nam','Hoa Kỳ','Châu Âu']]
stats={'date':DATE,'scope':'technology_only','schools':len(schools),'selected_work_records':len(works),'works_with_doi':sum(bool(w.get('doi')) for w in works),
 'regions':coverage,'primary_topic_counts':dict(counts),'all_topic_counts':dict(allcounts),'work_types':dict(collections.Counter(w['type'] for w in works)),
 'evidence_levels':dict(collections.Counter(w['evidence_level'] for w in works)),
 'bibliography_kept':len(bib),'bibliography_removed_or_uncertain':len(bio_source)-len(bib),
 'detailed_works_removed_or_uncertain':len(all_original)-len(works)}
rules=[
 'Giữ nghiên cứu có đóng góp về thuật toán, phần mềm, tính toán, an ninh mạng, điện tử–bán dẫn–viễn thông, robot–điều khiển, vật liệu, năng lượng, kỹ thuật môi trường, công nghệ sinh học/y sinh, xây dựng–giao thông hoặc công nghệ nông nghiệp.',
 'Loại nội dung thuần quản trị, kinh tế, xã hội, ngôn ngữ, giáo dục, tuyển sinh, sự kiện, lời mời hội nghị, quảng bá và khoa học cơ bản không thấy đóng góp công nghệ. Mục mơ hồ được loại khỏi kết quả chính.',
 'Nghiên cứu dùng công nghệ trong tài chính, giáo dục, pháp luật hoặc y tế vẫn được giữ nếu có phương pháp/hệ thống kỹ thuật cụ thể. Không loại chỉ vì lĩnh vực ứng dụng ngoài công nghệ.',
 'Các công trình lượng tử được xét theo đóng góp cụ thể: giữ thiết bị/hệ thống đo lường có triển khai, loại đề xuất chủ yếu nhằm kiểm tra vật lý cơ bản khi chưa thấy đóng góp công nghệ ứng dụng.',
 'Danh sách trường trong bản này chỉ gồm trường có công trình công nghệ phù hợp đã đọc nguồn. Danh bạ rộng và các đầu mối chưa xác minh được lưu trong hồ sơ làm việc, không đưa vào bộ kết quả chính.'
]
limitations=[
 'Báo cáo là mẫu công trình có nguồn của các trường được rà soát, chưa bao phủ mọi trường công nghệ trên thế giới hoặc toàn bộ 589 trường được nêu trong thống kê ban đầu. Không có tệp nguồn 589/5.021 trường để đối chiếu đầy đủ.',
 'Các mục chi tiết gồm bài công bố, bài đã chấp nhận, báo cáo, luận án, đồ án và dự án; có cột phân biệt. Phần lớn chỉ đọc abstract/bản tin cơ quan; một số mới xác nhận thông tin thư mục. Chưa thẩm định toàn văn hay khả năng tái lập.',
 'Nguồn chính thức có thể có lỗi, ngày hội nghị và ngày xuất bản có thể khác nhau; thông tin chưa chắc được giữ trong ghi chú. Không suy ra tác giả thuộc trường chỉ vì bài đăng trên tạp chí của trường; trường hợp chỉ có email hoặc kho cơ quan được nêu rõ.',
 'Phụ lục thư mục đã lọc liên quan công nghệ và gán nhóm theo tiêu đề, một số có đọc thêm abstract để xử lý trường hợp khó. Chưa kiểm tra đầy đủ nội dung từng bài. Không cộng trực tiếp số phụ lục với các mục chi tiết vì có thể trùng công trình.',
 'Một công trình có nhãn chính và nhiều nhãn phụ. Số mẫu không phản ánh tổng sản lượng, thế mạnh hoặc thứ hạng của trường. Trường đa ngành vẫn xuất hiện nếu có công trình công nghệ phù hợp.',
 'Website trường và cổng nghiên cứu được dẫn theo nguồn; không khẳng định đã rà soát mọi subdomain hay toàn bộ PDF. Không truy cập nội dung yêu cầu đăng nhập.'
]
meta={'date':DATE,'statistics':stats,'inclusion_exclusion_rules_vi':rules,'limitations_vi':limitations,'topic_taxonomy':taxonomy}
def dump(name,obj):(OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
dump('research_classified.json',{'metadata':meta,'schools':schools,'research_items':works})
dump('universities_tech_focus.json',{'metadata':meta,'technology_universities':[{k:v for k,v in s.items() if k!='items'} for s in schools]})
dump('bibliography_technology_only.json',{'metadata':{'date':DATE,'reading_depth':'titles_and_bibliographic_metadata_with_selected_abstract_checks','relevance':'manually_screened_title','topic_classification':'title_level_classification_reviewed','rules_vi':rules},'records':bib})
dump('statistics.json',stats)
# Audit stays in working evidence folder; excluded nontechnology rows are not shipped.
(ROOT/'technology_filter_audit.json').write_text(json.dumps(reviews,ensure_ascii=False,indent=2))

types={'journal_article':'Bài báo tạp chí','conference_paper':'Bài hội nghị','accepted_conference_paper':'Bài đã chấp nhận; chưa xác nhận xuất bản','book_chapter':'Chương sách/kỷ yếu','technical_report':'Báo cáo kỹ thuật','doctoral_dissertation':'Luận án tiến sĩ','doctoral_thesis':'Luận án tiến sĩ','student_graduation_project':'Đồ án tốt nghiệp','institutional_research_project':'Đề tài/dự án của cơ sở','funded_research_project':'Đề tài/dự án được tài trợ','research_project':'Dự án nghiên cứu','review_article':'Bài tổng quan'}
def tidy(v):
 t=str(v or '');t=re.sub(r'(?<=[^\W\d_])(?=20[12]\d\b)',' ',t);t=re.sub(r'(?<=[.,;])(?=[A-Za-zÀ-ỹ])',' ',t)
 for a,b in [('mạng5G','mạng 5G'),('thông1','thông 1'),('tác giảDUT','tác giả DUT'),('tác giảChungnam','tác giả Chungnam'),('thuộcHCMUTE','thuộc HCMUTE'),('khoVKU','kho VKU'),('báoVKU','báo VKU'),('websiteKhoa','website Khoa')]:t=t.replace(a,b)
 return t
def flat(v):return '' if v is None else '; '.join(map(str,v)) if isinstance(v,list) else v
headers=['ID','Trường','ISO2','Tên công trình','Năm','Loại','Chủ đề chính','Chủ đề phụ','Từ khóa chi tiết','Tóm tắt','DOI','URL nguồn','Nguồn đối chiếu','Mức chứng cứ','Ghi chú kiểm chứng']
data=[[w['record_id'],w['school'],w['country_code'],w['title'],w['year'],types.get(w['type'],w['type']),w['primary_topic_vi'],[taxonomy[c] for c in w['secondary_topic_codes']],w.get('topics',[]),tidy(w['summary_vi']),w.get('doi'),w['url'],w['source_urls'],w['evidence_level'],tidy(w.get('verification_note'))] for w in works]
with (OUT/'research_classified.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f);writer.writerow(headers);writer.writerows([[flat(v) for v in row] for row in data])
school_headers=['Trường','Tên ngắn','ISO2','Website','Cổng nghiên cứu','Số mục phù hợp','Chủ đề trong mẫu','Phạm vi']
school_data=[[s['name'],s['short_name'],s['country_code'],s['website'],[p['url'] for p in s['research_portals']],len(s['items']),sorted({taxonomy[c] for w in s['items'] for c in w['topic_codes']}),s['technology_evidence']] for s in schools]
with (OUT/'university_technology_websites.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f);writer.writerow(school_headers);writer.writerows([[flat(v) for v in row] for row in school_data])

wb=Workbook();wb.remove(wb.active)
def sheet(name,hs,ds,widths=None):
 ws=wb.create_sheet(name);ws.append(hs)
 for row in ds:
  clean=[]
  for v in row:
   v=flat(v)
   if isinstance(v,str):
    v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:32700]
    if v.startswith('='):v="'"+v
   clean.append(v)
  ws.append(clean)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False;ws.row_dimensions[1].height=32
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183B56');c.alignment=Alignment(wrap_text=True)
 for i in range(1,len(hs)+1):ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,25)
 for line in ws.iter_rows(min_row=2):
  for c in line:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str) and c.value.startswith(('https://','http://')) and '; ' not in c.value and len(c.value)<1800:c.hyperlink=c.value;c.style='Hyperlink'
 return ws
sheet('Huong_dan',['Mục','Nội dung'],[['Phạm vi','Chỉ giữ công trình liên quan công nghệ–kỹ thuật. Không coi danh sách này là đầy đủ toàn cầu.'],['Sử dụng','Cong_trinh: dữ liệu chi tiết; Truong_cong_nghe: website/cổng nghiên cứu; Chu_de: thống kê; Thu_muc_da_loc: phụ lục mới sàng lọc tiêu đề.']]+[[f'Tiêu chí {i+1}',t] for i,t in enumerate(rules)]+[[f'Giới hạn {i+1}',t] for i,t in enumerate(limitations)],{1:25,2:120})
sheet('Tong_quan',['Khu vực','Trường','Mục công trình'],[[r['region'],r['schools'],r['items']] for r in coverage],{1:30})
sheet('Chu_de',['Mã','Nhóm','Số mục: nhãn chính','Số mục: mọi nhãn'],[[k,v,counts[k],allcounts[k]] for k,v in taxonomy.items()],{1:30,2:60})
sheet('Cong_trinh',headers,data,{2:60,4:85,6:40,7:50,8:60,9:60,10:100,12:85,13:100,15:110})
sheet('Truong_cong_nghe',school_headers,school_data,{1:65,4:50,5:95,7:80,8:95})
sheet('Thu_muc_da_loc',['Trường','Tên công trình','Năm','Nhóm gợi ý','URL nguồn','URL công bố','DOI','Cách phân loại','Giới hạn'],[[r['school'],r['title'],r['year'],r['topics'],r['source_url'],r['publication_urls'],r.get('doi'),r['classification_method'],r['verification_note']] for r in bib],{2:100,4:70,5:85,6:85,8:100,9:100})
wb.save(OUT/'bao_cao_truong_cong_nghe.xlsx')

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.header.paragraphs[0].text='NGHIÊN CỨU CÔNG NGHỆ–KỸ THUẬT • ĐÃ LỌC NỘI DUNG NGOÀI PHẠM VI'
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
for st in ['Title','Heading 1','Heading 2','Heading 3']:doc.styles[st].font.color.rgb=RGBColor.from_string('183B56')
def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def para(t):return doc.add_paragraph(tidy(t))
doc.add_heading('Báo cáo phân loại nghiên cứu công nghệ–kỹ thuật',0)
para('Ngày thu thập: 30/10/2026. Bản đã lọc theo yêu cầu: chỉ giữ nội dung liên quan công nghệ.')
para(f'Kết quả chính: {len(works)} mục công trình tại {len(schools)} trường; {stats["works_with_doi"]} mục có DOI. Phụ lục giữ {len(bib)} bản ghi thư mục sau sàng lọc tiêu đề. Đây là mẫu được rà soát, chưa phải toàn bộ công trình của mọi trường công nghệ.')
doc.add_heading('1. Tiêu chí giữ và loại',1)
for r in rules:para(r)
doc.add_heading('2. Mức kiểm chứng và giới hạn',1)
for r in limitations:para(r)
doc.add_heading('3. Thống kê theo khu vực và chủ đề',1)
t=doc.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Khu vực','Trường','Mục công trình']):c.text=v
for r in coverage:
 for c,v in zip(t.add_row().cells,[r['region'],r['schools'],r['items']]):c.text=str(v)
t=doc.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Chủ đề','Số mục: nhãn chính','Số mục: mọi nhãn']):c.text=v
for k,v in taxonomy.items():
 for c,x in zip(t.add_row().cells,[v,counts[k],allcounts[k]]):c.text=str(x)
doc.add_heading('4. Chi tiết từng trường',1)
for num,s in enumerate(schools,1):
 doc.add_heading(f'{num}. {s["name"]}',2)
 p=doc.add_paragraph();p.add_run('Website: ');link(p,s['website'],s['website']);p.add_run(' | ISO2: '+s['country_code'])
 for port in s['research_portals']:
  p=doc.add_paragraph();p.add_run('Cổng tra cứu: ');link(p,tidy(port['title']),port['url'])
 for w in s['items']:
  doc.add_heading(w['record_id']+' • '+w['title'],3)
  para(f'Năm: {w["year"] if w["year"] else "chưa xác định"} | Loại: {types.get(w["type"],w["type"])}')
  para('Chủ đề chính: '+w['primary_topic_vi']+'. Chủ đề phụ: '+('; '.join(taxonomy[c] for c in w['secondary_topic_codes']) or 'Không')+'.')
  para(w['summary_vi']);para('Kiểm chứng: '+w.get('verification_note',''))
  p=doc.add_paragraph();link(p,'Nguồn nội dung',w['url'])
  if w.get('doi'):p.add_run(' | ');link(p,'DOI: '+w['doi'],'https://doi.org/'+w['doi'].removeprefix('https://doi.org/'))
  for u in w['source_urls']:
   if u!=w['url']:
    p=doc.add_paragraph();link(p,'Nguồn đối chiếu: '+u,u)
doc.add_heading('5. Tệp kèm theo',1)
para('Excel có các trang Cong_trinh, Truong_cong_nghe, Chu_de, Tong_quan và Thu_muc_da_loc. Dùng bộ lọc theo trường, quốc gia, nhóm công nghệ và loại công trình.')
para('research_classified.json/CSV chứa dữ liệu chi tiết; universities_tech_focus.json và university_technology_websites.csv chứa website các trường có công trình phù hợp; bibliography_technology_only.json là phụ lục thư mục chỉ mới sàng lọc tiêu đề.')
doc.save(OUT/'bao_cao_nghien_cuu_chi_tiet.docx')
with zipfile.ZipFile(ROOT/'bao_cao_truong_cong_nghe.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in OUT.iterdir():z.write(p,p.name)
print(json.dumps(stats,ensure_ascii=False,indent=2));print('Files',[(p.name,p.stat().st_size) for p in OUT.iterdir()])
