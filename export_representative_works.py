import sys,json,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'.deps'))
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from config import ROOT,DATE

OUT=ROOT/'deliverables'
data=json.loads((OUT/'research_classified.json').read_text())
selected=json.loads((ROOT/'representative_works_selected.json').read_text())
workmap={w['record_id']:w for w in data['research_items']}
translations={
 'USTH':'Chế tạo phức nickel làm xúc tác điện hóa tạo hydro',
 'HUST':'Dự đoán công dụng mới của thuốc bằng mạng tương đồng bệnh đa nguồn',
 'UET':'PRADA-QA: hỏi đáp sản phẩm bằng AI đa tác tử và truy xuất tri thức động',
 'HCMUT':'Lập lịch tác vụ siêu máy tính bằng học tăng cường sâu',
 'UIT':'Đánh giá độ bền vững và khả năng giải thích của LLM phát hiện lỗ hổng mã nguồn',
 'PTIT':'AUTOGENT: AI đa tác tử phát hiện lỗ hổng và cảnh báo an ninh mạng',
 'DUT':'Tay máy sáu bậc tự do tích hợp AI nhận dạng bao gói và tư thế sản phẩm',
 'VKU':'Phát hiện cháy trong nhà bằng thị giác máy tính kết hợp cảm biến IoT',
 'HCMUTE':'Điều khiển robot di động bám quỹ đạo và tránh vật cản động',
 'VLUTE':'AI có thể giải thích hỗ trợ nhận diện Alzheimer từ ảnh MRI',
 'MTA':'Đánh giá tính bền vững khi chế tạo hợp kim nickel bằng đắp dây hồ quang',
 'KMA':'Tính toán bảo mật cho hệ gợi ý có đảm bảo quyền riêng tư',
 'UTC':'Vữa dùng phụ phẩm công nghiệp cho mặt đường bán mềm',
 'HUMG':'Giám sát dao động cầu dây văng theo thời gian thực bằng IoT và đám mây',
 'HUCE':'Mô hình thị giác–ngôn ngữ nhỏ gọn để phân loại tình huống cháy trong nhà',
 'FPT':'AI và học máy cá nhân hóa trải nghiệm học tập',
 'Phenikaa':'Định vị robot tự hành bằng UWB trong nhà máy thông minh',
 'MIT':'Thiết kế robot nhảy bằng AI tạo sinh và mô phỏng vật lý',
 'Caltech':'Mạng nơ-ron bằng DNA có thể học từ các ví dụ hóa học',
 'Georgia Tech':'SAIL: tăng tốc thao tác của robot học từ động tác mẫu của người',
 'CMU':'Khảo sát cách vận dụng kiểm thử đối kháng an ninh mạng vào AI tạo sinh',
 'Purdue':'Phân tích năng lượng và khí hậu của sản xuất thép dùng hydro',
 'RPI':'Tách màng tinh thể siêu mỏng cho cảm biến hồng ngoại không cần làm lạnh',
 'Virginia Tech':'Kết hợp kẹp cứng và kẹp mềm để robot thao tác đồ vật linh hoạt',
 'Texas A&M':'Pin hữu cơ hai ion hoạt động ở nhiệt độ thấp và có khả năng gia cường kết cấu',
 'Stevens':'Nhận dạng hoạt động người từ cảm biến bằng học tự giám sát',
 'IIT':'Thiết kế chất điện phân rắn oxyhalide dẫn ion lithium cho pin thể rắn',
 'ETH Zurich':'RAMBO: điều khiển toàn thân để robot vừa di chuyển vừa thao tác đồ vật',
 'EPFL':'Thiết kế lớp tiếp xúc điện tử cho pin mặt trời perovskite đảo cấu trúc',
 'Imperial':'Học chuyển giao để tăng tốc tính nhiệt hóa học trong mô phỏng cháy',
 'TU Delft':'DRONE-RL: UAV tự điều hướng khi môi trường thay đổi',
 'TU/e':'Cảm biến nhiệt độ InP kết nối Si-CMOS cho mạch quang tích hợp',
 'RWTH Aachen':'Điều khiển tối ưu động cơ hydro đánh lửa cưỡng bức',
 'TUM':'Thiết kế vật liệu bằng học máy dựa trên đặc tính vận chuyển điện tích',
 'TU Berlin':'Học tăng cường sâu cho lập kế hoạch chuyển động robot công nghiệp cộng tác',
 'KTH':'AttentionGuard: phát hiện dữ liệu giả mạo trong đoàn xe kết nối bằng Transformer',
 'Chalmers':'Ước lượng nguy cơ mạ lithium bằng học máy để hỗ trợ sạc pin nhanh',
 'Aalto':'Làm lạnh nhanh mạch lượng tử bằng xuyên hầm hỗ trợ photon',
 'Politecnico di Milano':'Sản xuất hydro từ nhiệt phân methane kết hợp thu giữ CO₂',
 'Politecnico di Torino':'Điều khiển có ràng buộc an toàn cho robot phẫu thuật dạng liên tục',
 'TU Wien':'Advanced Bioenergy Lab Austria: hạ tầng thí nghiệm nhiên liệu bền vững từ sinh khối',
}
types={'journal_article':'Bài báo tạp chí','conference_paper':'Bài hội nghị','book_chapter':'Chương sách/kỷ yếu',
 'accepted_conference_paper':'Bài đã chấp nhận; chưa xác nhận xuất bản','doctoral_thesis':'Luận án tiến sĩ',
 'funded_research_project':'Đề tài/dự án được tài trợ','institutional_research_project':'Đề tài cấp cơ sở',
 'technical_report':'Báo cáo kỹ thuật','research_project':'Dự án nghiên cứu'}
records=[]
for r in selected:
 w=workmap[r['record_id']]
 records.append({**r,'title_vi':translations[r['school']],'region':w['region'],'country_code':w['country_code'],
  'type_vi':types[r['type']],'topics_vi':w['topic_labels_vi'],'doi':w.get('doi'),
  'verification_note':w['verification_note'],'source_urls':w['source_urls']})
assert len(records)==len(set(r['school'] for r in records))==41
metadata={'date':DATE,'selection_note_vi':'Mỗi trường chọn một công trình minh họa trong tập đã rà soát; không xếp hạng công trình tốt nhất, trích dẫn nhiều nhất hoặc toàn bộ thế mạnh của trường. Tên tiếng Việt là diễn giải ngắn; tên gốc được giữ riêng.',
 'reading_note_vi':'Bài công bố, bài đã chấp nhận, luận án, báo cáo và dự án được ghi loại riêng; xem ghi chú kiểm chứng để biết mức đọc nguồn.'}
(OUT/'de_tai_tieu_bieu_41_truong.json').write_text(json.dumps({'metadata':metadata,'records':records},ensure_ascii=False,indent=2))
headers=['Trường','Khu vực','ISO2','Tên diễn giải tiếng Việt','Tên gốc','Năm','Loại công trình','Nhóm công nghệ','Nội dung nghiên cứu','DOI','URL nguồn','Nguồn đối chiếu','Ghi chú kiểm chứng']
rows=[[r['school'],r['region'],r['country_code'],r['title_vi'],r['title'],r['year'],r['type_vi'],'; '.join(r['topics_vi']),r['summary'],r.get('doi'),r['url'],'; '.join(r['source_urls']),r['verification_note']] for r in records]
def fill(ws):
 ws.append(headers)
 for row in rows:ws.append(row)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
 widths=[25,20,8,65,80,10,35,50,90,40,75,85,100]
 for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183B56');cell.alignment=Alignment(wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
  row[10].hyperlink=row[10].value;row[10].style='Hyperlink'
wb=Workbook();wb.active.title='De_tai_tieu_bieu';fill(wb.active)
note=wb.create_sheet('Pham_vi');note.append(['Mục','Nội dung'])
for k,v in metadata.items():note.append([k,v])
note.column_dimensions['A'].width=30;note.column_dimensions['B'].width=120
for row in note:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(OUT/'de_tai_tieu_bieu_41_truong.xlsx')
main=load_workbook(OUT/'bao_cao_truong_cong_nghe.xlsx')
if 'De_tai_tieu_bieu' in main:del main['De_tai_tieu_bieu']
fill(main.create_sheet('De_tai_tieu_bieu',2));main.save(OUT/'bao_cao_truong_cong_nghe.xlsx')
with zipfile.ZipFile(ROOT/'bao_cao_truong_cong_nghe.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in OUT.iterdir():z.write(p,p.name)
print('Exported41representativeworks; mainworkbookupdated')
