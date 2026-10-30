import sys,json,re,csv,zipfile,collections
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'.deps'))
from openpyxl import load_workbook
from docx import Document
from ai_config import ROOT,AI_TOPICS

out=ROOT/'ai_deliverables';d=json.loads((out/'ai_research_catalog.json').read_text());stats=d['metadata']['statistics']
schools=d['schools'];works=d['unique_works'];records=d['institution_research_records']
base=json.loads((ROOT/'deliverables/research_classified.json').read_text())
assert len(schools)==len({s['short_name'] for s in schools})==41
assert {s['short_name'] for s in schools}=={s['short_name'] for s in base['schools']}
assert len(works)==len({w['work_key'] for w in works})==stats['unique_works']
assert len(records)==len({r['association_id'] for r in records})==stats['institution_work_records']
assert sum(len(s['items']) for s in schools)==len(records)
assert min(len(s['items']) for s in schools)>=2
assert sum(stats['primary_topic_counts'].values())==len(works)
assert sum(x['institution_work_records'] for x in stats['regional_coverage'])==len(records)
for r in records:
 for field in ['title','year','type','status','school_attribution','url','summary_vi','problem_vi','ai_method_vi','reading_depth','verification_note','source_urls','ai_topic_codes','application_domains']:
  assert r.get(field), (r['title'],field)
 assert set(r['ai_topic_codes'])<=set(AI_TOPICS),r['title']
 assert r['primary_ai_topic_code'] in r['ai_topic_codes'],r['title']
 assert 2024<=r['year']<=2026,(r['title'],r['year'])
 assert not r['full_text_read']
 if r.get('doi'):assert re.match(r'^10\.\d{4,9}/\S+$',r['doi']),r['doi']
 assert all(u.startswith(('https://','http://')) for u in r['source_urls']),r['title']
 assert r['work_id'] in {w['work_id'] for w in works}
for w in works:
 matches=[r for r in records if r['work_id']==w['work_id']]
 assert set(w['schools'])=={r['school'] for r in matches}
 assert len(w['association_ids'])==len(matches)
# Important real source/date discrepancies must remain visible.
load=next(w for w in works if w['title'].startswith('Enhanced Monthly Load'))
assert load['year']==2026 and '2025' in load['doi']
sbert=next(w for w in works if w['title'].startswith('Fine-Tuned Sentence-BERT'))
assert sbert['year']==2026 and '2018' in sbert['verification_note'] and not sbert.get('doi')
ecg=next(w for w in works if w['title'].startswith('MULTI-LEVEL STACKING'))
assert ecg['alternative_years']==[2024]
shared=[w for w in works if len(w['schools'])>1]
assert len(shared)==2
with (out/'ai_research_catalog.csv').open(encoding='utf-8-sig') as f:assert len(list(csv.reader(f)))==len(works)+1
wb=load_workbook(out/'bao_cao_nghien_cuu_AI.xlsx',read_only=True,data_only=False)
assert wb['Cong_trinh_AI'].max_row==len(works)+1
assert wb['Theo_truong'].max_row==42
assert wb['Ban_ghi_truong'].max_row==len(records)+1
assert wb['Chu_de_AI'].max_row==len(AI_TOPICS)+1
assert not any(c.data_type=='f' for ws in wb for row in ws for c in row)
wb.close()
doc=Document(out/'bao_cao_nghien_cuu_AI.docx')
assert len([p for p in doc.paragraphs if p.style.name=='Heading 3'])==len(records)
with zipfile.ZipFile(ROOT/'bao_cao_AI_cac_truong.zip') as z:
 assert z.testzip() is None
 assert set(z.namelist())=={p.name for p in out.iterdir()}
 for n in z.namelist():
  if n.endswith('.json'):json.loads(z.read(n))
result={'status':'PASS','schools':len(schools),'unique_ai_works':len(works),'institution_work_records':len(records),
 'new_vs_previous_detailed_ai_subset':stats['new_unique_works_vs_previous_detailed_ai_subset'],
 'topic_groups':len(AI_TOPICS),'shared_works_deduplicated':len(shared),'doi_count':stats['unique_works_with_doi'],
 'no_missing_school':True,'xlsx_docx_json_csv_zip_readable':True,'full_text_review_claimed':False,
 'scope':'41 previously listed institutions, selected works only'}
(ROOT/'ai_validation_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
