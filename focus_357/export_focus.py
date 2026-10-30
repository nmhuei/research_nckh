import collections
import csv
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import unquote

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
sys.path.insert(0, str(ROOT / '.deps'))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

DATE = '2026-10-30'
OUT = BASE / 'deliverables'
OUT.mkdir(exist_ok=True)
GROUPS = {
    'G3': 'AI tiết kiệm tài nguyên và chạy trên thiết bị',
    'G5': 'AI cho robot và hệ thống tự hành',
    'G7': 'AI tối ưu năng lượng và hệ thống kỹ thuật',
}
scope = json.loads((BASE / 'school_scope.json').read_text())
school_map = {s['short_name']: s for s in scope}
order = {s['short_name']: i for i, s in enumerate(scope)}
old = json.loads((ROOT / 'ai_deliverables/ai_research_catalog.json').read_text())['unique_works']
screening = json.loads((BASE / 'bibliography_screening.json').read_text())['records']
techdoc = json.loads((BASE / 'parts/technology_intersection.json').read_text())

def normtitle(value):
    value = unicodedata.normalize('NFKD', str(value)).casefold()
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in value if not unicodedata.combining(c)))

def normdoi(value):
    if not value:
        return None
    value = unquote(str(value)).strip()
    value = re.sub(r'^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)', '', value, flags=re.I)
    return value.lower().rstrip(' .') or None

def unique(values):
    return list(dict.fromkeys(v for v in values if v is not None and v != ''))

def arr(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]

def clean(value):
    value = str(value or '')
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)

def flat(value):
    if value is None:
        return ''
    if isinstance(value, (list, tuple)):
        return '; '.join(flat(x) for x in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value

raw_records = []
audits = []
for name in ['group3', 'group5', 'group7', 'root_verified', 'technology_additions']:
    part = {'items': techdoc.get('additional_items', [])} if name == 'technology_additions' else json.loads((BASE / 'parts' / (name + '.json')).read_text())
    for raw in part['items']:
        r = dict(raw)
        r['part'] = name
        r['doi'] = normdoi(r.get('doi'))
        r['schools'] = unique(arr(r.get('schools')))
        r['group_codes'] = sorted(unique(arr(r.get('group_codes'))), key=lambda x: list(GROUPS).index(x))
        r['source_urls'] = unique([r.get('url')] + arr(r.get('source_urls')))
        r['checked_on'] = DATE
        r['full_text_read'] = bool(r.get('full_text_read', False))
        r['subtopics_vi'] = unique(arr(r.get('subtopics_vi')))
        r['ai_methods'] = unique(arr(r.get('ai_methods')))
        raw_records.append(r)
    for a in part.get('school_audit', []):
        audits.append(dict(a, group_code=part['group_code']))

# Preserve all evidence when school/group researchers find the same work.
clusters = []
by_doi = {}
by_title = {}
for r in raw_records:
    d, t = r['doi'], normtitle(r['title'])
    hits = unique([by_doi.get(d) if d else None, by_title.get(t)])
    if hits:
        idx = hits[0]
        for other in hits[1:]:
            if other != idx:
                clusters[idx].extend(clusters[other])
                clusters[other] = []
                by_doi = {k: idx if v == other else v for k, v in by_doi.items()}
                by_title = {k: idx if v == other else v for k, v in by_title.items()}
    else:
        idx = len(clusters)
        clusters.append([])
    clusters[idx].append(r)
    if d:
        by_doi[d] = idx
    by_title[t] = idx

def evidence_rank(r):
    depth = r.get('reading_depth', '').lower()
    if r['part'] == 'root_verified':
        return 100
    if 'abstract' in depth or 'methods' in depth or 'selected_sections' in depth:
        return 50
    if 'bibliograph' in depth or 'title_only' in depth:
        return 5
    return 25

works = []
weak_records = []
for cluster in clusters:
    if not cluster:
        continue
    cluster.sort(key=evidence_rank, reverse=True)
    w = dict(cluster[0])
    w['schools'] = sorted(unique(s for r in cluster for s in r['schools']), key=order.get)
    w['group_codes'] = sorted(unique(g for r in cluster for g in r['group_codes']), key=list(GROUPS).index)
    for key in ['source_urls', 'subtopics_vi', 'ai_methods']:
        w[key] = unique(x for r in cluster for x in r.get(key, []))
    w['doi'] = next((r['doi'] for r in cluster if r.get('doi')), None)
    w['evidence_records'] = cluster
    w['group_names_vi'] = [GROUPS[g] for g in w['group_codes']]
    w['regions'] = unique(school_map[s]['region'] for s in w['schools'])
    w['country_codes'] = unique(school_map[s]['country_code'] for s in w['schools'])
    w['source_years'] = unique(r.get('year') for r in cluster)
    w['institution_evidence_vi'] = '\n'.join(unique(r.get('institution_evidence_vi') for r in cluster))
    w['verification_notes_vi'] = unique(r.get('verification_note_vi') for r in cluster)
    # Efficient operation of a general computer system is G7; G3 needs a
    # demonstrated contribution to the cost/deployment of AI itself.
    g3_boundary_titles = {
        normtitle('Toward Sustainable and Cost-Efficient HPC Systems: A Deep Reinforcement Learning Job Scheduling Approach'),
        normtitle('Fine-Tuned Sentence-BERT for HPC Job Outcome Prediction via Textual Feature Embedding'),
        normtitle('FedCCL: Federated Clustered Continual Learning Framework for Privacy-Focused Energy Forecasting'),
    }
    if normtitle(w['title']) in g3_boundary_titles and 'G3' in w['group_codes']:
        w['group_codes'].remove('G3')
        w['group_names_vi'] = [GROUPS[g] for g in w['group_codes']]
        w['verification_notes_vi'].append('Rà soát cuối: giữ G7; chưa đủ chứng cứ giảm tài nguyên của mô hình AI hoặc triển khai AI trên thiết bị để gắn G3. Học liên kết riêng tư không tự đồng nghĩa tiết kiệm tài nguyên.')
    additions = {
        normtitle('CoFANN: A collaborative framework for accelerating DNN inference in drone-based agricultural monitoring systems'): ('G5', 'G5 là thành phần nhận thức học sâu của hệ giám sát bằng drone; kết quả đánh giá suy luận trên dữ liệu, chưa xác nhận điều khiển bay hay giảm năng lượng bay.'),
        normtitle('Cooperative Edge Inferences With Online Learning'): ('G7', 'G7 là tối ưu năng lượng và độ trễ của hạ tầng suy luận biên, không phải năng lượng động cơ robot.'),
        normtitle('Boosting the Performance of Deep Reinforcement Learning for Energy Management Systems using Behavior Cloning from Linear Programming Solutions'): ('G3', 'G3 là giảm số bước huấn luyện chính sách nhờ behavior cloning; không khẳng định chạy trên vi điều khiển.'),
    }
    if normtitle(w['title']) in additions:
        code, note = additions[normtitle(w['title'])]
        w['group_codes'] = sorted(unique(w['group_codes'] + [code]), key=list(GROUPS).index)
        w['group_names_vi'] = [GROUPS[g] for g in w['group_codes']]
        w['verification_notes_vi'].append('Rà soát giao công nghệ: ' + note)
    if len(w['source_years']) > 1:
        w['verification_notes_vi'].append('Các nguồn có năm khác nhau: ' + ', '.join(map(str, w['source_years'])) + '. Xem từng bằng chứng; ưu tiên năm trong bản ghi được chọn.')
    depth = w.get('reading_depth', '').lower()
    if all(('bibliograph' in r.get('reading_depth', '').lower() or 'title_only' in r.get('reading_depth', '').lower()) and 'abstract' not in r.get('reading_depth', '').lower() and 'content_reviewed' not in r.get('reading_depth', '').lower() for r in cluster):
        weak_records.append(w)
        continue
    w['evidence_level'] = 'abstract_or_content_reviewed' if ('abstract' in depth or 'methods' in depth or 'content_reviewed' in depth or 'selected_sections' in depth) else 'official_metadata_or_project_registry'
    works.append(w)

works.sort(key=lambda w: (min(order[s] for s in w['schools']), -(w.get('year') or 0), w['title']))
for i, w in enumerate(works, 1):
    w['work_id'] = f'F{i:03}'

work_by_title = {normtitle(w['title']): w for w in works}
work_by_doi = {w['doi']: w for w in works if w.get('doi')}
techs = []
for raw in techdoc['technologies']:
    t = dict(raw)
    t['work_ids_by_group'] = {}
    t['unmatched_titles'] = []
    for g, titles in t['work_title_evidence'].items():
        t['work_ids_by_group'][g] = []
        for title in titles:
            matched = work_by_title.get(normtitle(title))
            if matched:
                t['work_ids_by_group'][g].append(matched['work_id'])
            else:
                t['unmatched_titles'].append(title)
    t['all_work_ids'] = unique(i for ids in t['work_ids_by_group'].values() for i in ids)
    t['direct_triple_work_ids'] = [w['work_id'] for w in works if w['work_id'] in t['all_work_ids'] and all(g in w['group_codes'] for g in GROUPS)]
    # Add the direct robot-ROM evidence to the engineering-model leg explicitly.
    if t['technology_id'] == 'TI01':
        robotrom = work_by_title[normtitle('Structure-Preserving Model Order Reduction of Slender Soft Robots via Autoencoder-Parameterized Strain')]
        t['work_ids_by_group']['G7'] = unique(t['work_ids_by_group']['G7'] + [robotrom['work_id']])
    techs.append(t)
strong_techs = [t for t in techs if t.get('strong_filter')]
support_ids = {i for t in strong_techs for i in t['all_work_ids']}
support_works = [w for w in works if w['work_id'] in support_ids]
leads = []
for r in screening:
    if not r['group_codes']:
        continue
    match = work_by_doi.get(normdoi(r.get('doi'))) or work_by_title.get(normtitle(r['title']))
    r['matched_work_id'] = match['work_id'] if match else None
    if not match:
        leads.append(dict(r, lead_sources=[r.get('source_url')] + r.get('publication_urls', [])))
for w in weak_records:
    if normtitle(w['title']) in {normtitle(r['title']) for r in leads}:
        continue
    leads.append(dict(title=w['title'], schools=w['schools'], school='; '.join(w['schools']), year=w.get('year'), group_codes=w['group_codes'], doi=w.get('doi'), lead_sources=w['source_urls'], verification_note='Chỉ xác minh thư mục; không nhập vào số công trình có nội dung/metadata chuyên biệt được kiểm tra.'))
for i, r in enumerate(leads, 1):
    r['lead_id'] = f'L{i:03}'

old_titles = {normtitle(w['title']) for w in old}
old_dois = {normdoi(w.get('doi')) for w in old if w.get('doi')}
new_count = sum(normtitle(w['title']) not in old_titles and w.get('doi') not in old_dois for w in works)
for s in scope:
    s['work_ids'] = [w['work_id'] for w in works if s['short_name'] in w['schools']]
    s['group_counts'] = {g: sum(s['short_name'] in w['schools'] and g in w['group_codes'] for w in works) for g in GROUPS}

stats = {
    'date': DATE,
    'schools_in_search_scope': len(scope),
    'schools_with_works': sum(bool(s['work_ids']) for s in scope),
    'unique_works': len(works),
    'new_vs_previous_87_ai_works': new_count,
    'school_work_associations': sum(len(w['schools']) for w in works),
    'group_counts': {g: sum(g in w['group_codes'] for w in works) for g in GROUPS},
    'group_school_counts': {g: sum(s['group_counts'][g] > 0 for s in scope) for g in GROUPS},
    'works_with_doi': sum(bool(w.get('doi')) for w in works),
    'multi_group_works': sum(len(w['group_codes']) > 1 for w in works),
    'direct_triple_group_works': sum(all(g in w['group_codes'] for g in GROUPS) for w in works),
    'technology_families_retained_across_three_groups': len(strong_techs),
    'technology_supporting_works': len(support_works),
    'bibliographic_leads_not_in_main_catalogue': len(leads),
    'previous_bibliography_entries_screened': len(screening),
    'school_group_search_audit_records': len(audits),
    'status_counts': dict(collections.Counter(w['status'] for w in works)),
    'evidence_level_counts': dict(collections.Counter(w['evidence_level'] for w in works)),
    'year_counts': dict(sorted(collections.Counter(str(w.get('year')) for w in works).items())),
    'regional_coverage': [{
        'region': region,
        'schools_in_scope': sum(s['region'] == region for s in scope),
        'schools_with_works': sum(s['region'] == region and bool(s['work_ids']) for s in scope),
        'unique_works_in_region': sum(region in w['regions'] for w in works),
    } for region in ['Việt Nam', 'Hoa Kỳ', 'Châu Âu']],
}
rules = [
    'G3: mô hình AI gọn, chắt lọc/nén/lượng tử hóa/cắt tỉa, suy luận hiệu quả, tăng tốc phần cứng, học liên kết tiết kiệm tài nguyên hoặc triển khai trên thiết bị biên. Không giữ một bài chỉ vì có chữ efficient nếu không thấy đóng góp phù hợp.',
    'G5: AI cho thiết kế, nhận thức, lập kế hoạch hoặc điều khiển robot và hệ tự hành. Không gán mọi bài học tăng cường vào robot; lập lịch HPC thuộc G7. Điều khiển cổ điển hoặc tối ưu metaheuristic thuần túy không được tự coi là robot học.',
    'G7: AI cho dự báo, vận hành, giám sát hoặc tối ưu năng lượng và hệ thống kỹ thuật, gồm sản xuất, bảo trì, công trình và hạ tầng tính toán. Không nhập nghiên cứu thuần hóa học/vật liệu/y tế vào G7 khi không có mục tiêu hệ thống kỹ thuật thuộc phạm vi.',
    'Dùng nguồn sơ cấp: cổng trường/nhóm tác giả, kho cơ quan, nhà xuất bản hoặc bản gốc. Quan hệ trường được ghi theo từng bằng chứng; tạp chí thuộc một trường không tự chứng minh các tác giả thuộc trường đó.',
    'Gộp DOI hoặc tiêu đề chuẩn hóa, giữ tất cả nhãn và bằng chứng. Một công trình hợp tác hoặc thuộc nhiều nhóm chỉ tính một lần ở tổng công trình duy nhất.',
]
limits = [
    'Phạm vi là 41 trường có tên cụ thể trong danh sách hiện có, không phải toàn bộ các trường trên thế giới hay danh sách 589 trường chưa được cung cấp dưới dạng tệp.',
    'Đã chạy truy vấn theo từng trường và từng nhóm, rà danh mục AI cũ cùng 287 bản ghi thư mục công nghệ. Tìm kiếm web có giới hạn truy cập và lập chỉ mục: đây là danh mục tìm được, chưa đủ căn cứ khẳng định đã thu toàn bộ sản lượng nghiên cứu của mỗi trường.',
    'Ưu tiên công trình 2024–2026, đến ngày 30/10/2026; giữ một số công trình nền tảng cũ hơn khi tìm thấy và phù hợp. Không áp dụng giới hạn năm tuyệt đối để loại bài có liên quan.',
    'Mức đọc và bằng chứng được ghi theo từng mục. Nội dung có thể chỉ là abstract, phần chọn lọc hoặc metadata/đăng ký đề tài chính thức; không khẳng định đã đọc hoặc tái lập toàn văn.',
    'Đề tài được duyệt, dự án đang triển khai, bản tiền công bố, bài được chấp nhận và bài đã công bố được giữ đúng trạng thái nguồn. Mục tiêu dự án không được trình bày thành kết quả đã đạt.',
    'Đầu mối chỉ có thư mục/tiêu đề được đặt trong phụ lục riêng và không cộng vào tổng công trình chính. Không tìm thấy ở một truy vấn không có nghĩa trường không có nghiên cứu ở nhóm đó.',
    'Số mục trong danh mục phụ thuộc phạm vi nguồn và cách tìm; không dùng để xếp hạng trường hoặc tính tốc độ tăng của lĩnh vực. Tổng số mục từng nhóm có thể lớn hơn tổng duy nhất vì nhãn giao nhau.',
]
catalogue = {'metadata': {'date': DATE, 'groups': GROUPS, 'statistics': stats, 'inclusion_rules_vi': rules, 'limitations_vi': limits}, 'schools': scope, 'works': works, 'technology_filter': {'scope_note_vi': techdoc['scope_note_vi'], 'technologies': techs}, 'bibliographic_leads': leads, 'search_audit': audits, 'previous_bibliography_screening': screening}
(OUT / 'catalogue_AI_3_5_7.json').write_text(json.dumps(catalogue, ensure_ascii=False, indent=2))
(OUT / 'statistics.json').write_text(json.dumps(stats, ensure_ascii=False, indent=2))
(OUT / 'cong_nghe_giao_3_nhom.json').write_text(json.dumps({'date': DATE, 'scope_note_vi': techdoc['scope_note_vi'], 'selection_note_vi': 'Giữ họ công nghệ có bằng chứng ở ba nhóm; ghi riêng công trình thực sự có đủ ba nhãn và công nghệ có bằng chứng qua nhiều công trình. G7 được chú thích rõ là năng lượng suy luận, năng lượng vận hành hoặc tối ưu mô hình/hệ thống kỹ thuật.', 'technologies': strong_techs, 'supporting_works': support_works, 'not_retained': [t for t in techs if not t.get('strong_filter')]}, ensure_ascii=False, indent=2))

headers = ['ID', 'Trường', 'Khu vực', 'Tên công trình gốc', 'Năm', 'Nhóm', 'Chủ đề cụ thể', 'Loại công trình', 'Trạng thái nguồn', 'Bài toán', 'Thuật toán', 'Phương pháp', 'Đóng góp / mục tiêu', 'Ứng dụng', 'DOI', 'URL chính', 'Nguồn đối chiếu', 'Bằng chứng trường', 'Mức đọc', 'Mức bằng chứng', 'Ghi chú kiểm chứng', 'Code', 'Dữ liệu', 'Ngày kiểm tra']
def workrow(w):
    return [w['work_id'], w['schools'], w['regions'], w['title'], w.get('year'), w['group_codes'], w['subtopics_vi'], w.get('type'), w['status'], w.get('problem_vi'), w['ai_methods'], w.get('method_vi'), w.get('contribution_vi'), w.get('application_vi'), w.get('doi'), w['url'], w['source_urls'], w['institution_evidence_vi'], w.get('reading_depth'), w['evidence_level'], w['verification_notes_vi'], w.get('code_url'), w.get('data_url'), DATE]
rows = [workrow(w) for w in works]
with (OUT / 'catalogue_AI_3_5_7.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([flat(v) for v in row])

wb = Workbook()
wb.remove(wb.active)
def sheet(name, cols, data, widths=None):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for row in data:
        cleanrow = []
        for value in row:
            value = flat(value)
            if isinstance(value, str):
                value = clean(value)[:32700]
                if value.startswith('='):
                    value = "'" + value
            cleanrow.append(value)
        ws.append(cleanrow)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='163A52')
        c.font = Font(bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 36
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 27)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(c.value, str) and c.value.startswith(('https://', 'http://')) and '; ' not in c.value and '\n' not in c.value:
                c.hyperlink = c.value
                c.style = 'Hyperlink'
    return ws

sheet('Huong_dan', ['Mục', 'Nội dung'], [['Ngày tra cứu', DATE], ['Phạm vi', '41 trường; ba nhóm G3, G5, G7.'], ['Số liệu', f'{len(works)} công trình chính; {len(leads)} đầu mối thư mục chưa đủ kiểm tra nội dung.'], ['Cách dùng', 'Loc theo G3/G5/G7 hoặc trường. ID công trình giống nhau giữa các sheet biểu thị cùng một công trình. Kiem_chung_truong chứa nhật ký truy vấn của từng trường và nhóm.']] + [[f'Tiêu chí {i+1}', t] for i, t in enumerate(rules)] + [[f'Giới hạn {i+1}', t] for i, t in enumerate(limits)], {1:25, 2:140})
sheet('Tong_quan', ['Mã nhóm', 'Nhóm nghiên cứu', 'Công trình có nhãn', 'Trường có công trình'], [[g, GROUPS[g], stats['group_counts'][g], stats['group_school_counts'][g]] for g in GROUPS], {1:15, 2:70, 3:30, 4:30})
techheaders = ['ID công nghệ', 'Họ công nghệ', 'Công nghệ cụ thể', 'Vai trò G3: AI gọn/thiết bị', 'Vai trò G5: robot', 'Vai trò G7: năng lượng/kỹ thuật', 'Cơ chế liên hệ', 'ID bằng chứng G3', 'ID bằng chứng G5', 'ID bằng chứng G7', 'ID công trình đủ ba nhãn', 'Mức kiểm chứng', 'Giới hạn / mức triển khai', 'Nguồn gốc']
def techrow(t):
    return [t['technology_id'], t['name_vi'], t['concrete_technologies'], t['role_g3_vi'], t['role_g5_vi'], t['role_g7_vi'], t['mechanism_vi'], t['work_ids_by_group'].get('G3'), t['work_ids_by_group'].get('G5'), t['work_ids_by_group'].get('G7'), t['direct_triple_work_ids'], t['evidence_note_vi'], t['maturity_note_vi'], t['source_urls']]
techwidths = {1:18, 2:65, 3:80, 4:100, 5:100, 6:110, 7:120, 8:35, 9:35, 10:35, 11:35, 12:125, 13:120, 14:110}
sheet('Cong_nghe_giao_3_nhom', techheaders, [techrow(t) for t in strong_techs], techwidths)
widths = {2:42, 3:25, 4:95, 6:18, 7:65, 8:30, 9:38, 10:90, 11:60, 12:100, 13:100, 14:65, 15:40, 16:90, 17:100, 18:100, 19:45, 20:50, 21:120, 22:70, 23:70}
sheet('Tat_ca_cong_trinh', headers, rows, widths)
for g in GROUPS:
    sheet({'G3':'G3_AI_nhe_va_thiet_bi', 'G5':'G5_Robot_tu_hanh', 'G7':'G7_Nang_luong_ky_thuat'}[g], headers, [workrow(w) for w in works if g in w['group_codes']], widths)
sheet('Theo_truong', ['Trường', 'Tên đầy đủ', 'Khu vực', 'Website', 'G3', 'G5', 'G7', 'Tổng duy nhất', 'ID công trình', 'Ghi chú'], [[s['short_name'], s['name'], s['region'], s['website'], s['group_counts']['G3'], s['group_counts']['G5'], s['group_counts']['G7'], len(s['work_ids']), s['work_ids'], 'Chưa tìm được mục đủ nguồn trong phạm vi tìm kiếm.' if not s['work_ids'] else 'Số mẫu không phải tổng sản lượng của trường.'] for s in scope], {2:70, 4:55, 9:90, 10:80})
sheet('Chu_de_cu_the', ['Nhóm', 'Chủ đề gắn nhãn', 'Số công trình', 'ID'], [[g, topic, len(ids), ids] for g in GROUPS for topic, ids in sorted({t:[w['work_id'] for w in works if g in w['group_codes'] and t in w['subtopics_vi']] for t in {t for w in works if g in w['group_codes'] for t in w['subtopics_vi']}}.items())], {2:85, 4:90})
sheet('Dau_moi_can_doi_chieu', ['ID đầu mối', 'Trường', 'Tên gốc', 'Năm theo danh mục', 'Nhóm dự kiến', 'Nguồn', 'DOI theo danh mục', 'Ghi chú'], [[r['lead_id'], r.get('school', r.get('schools')), r['title'], r.get('year'), r['group_codes'], unique(r.get('lead_sources', [])), r.get('doi'), r.get('verification_note', r.get('screening_note_vi'))] for r in leads], {2:40, 3:100, 6:110, 8:120})
sheet('Kiem_chung_truong', ['Nhóm', 'Trường', 'Truy vấn đã chạy', 'Nguồn truy cập / đối chiếu', 'Kết quả và giới hạn'], [[a['group_code'], a['school'], a.get('queries', []), a.get('source_urls', []), a.get('result_note_vi')] for a in audits], {3:115, 4:100, 5:125})
sheet('Ra_soat_287_thu_muc', ['ID thư mục', 'Trường', 'Tiêu đề', 'Nhóm dự kiến', 'Sàng lọc', 'ID đã đối chiếu', 'Nguồn', 'Ghi chú'], [[r['bibliography_id'], r['school'], r['title'], r['group_codes'], r['focus_screening'], r.get('matched_work_id'), r.get('source_url'), r['screening_note_vi']] for r in screening], {3:105, 5:65, 7:95, 8:110})
wb.save(OUT / 'bao_cao_AI_nhom_3_5_7.xlsx')
fullwb = wb
wb = Workbook()
wb.remove(wb.active)
sheet('Huong_dan', ['Mục', 'Nội dung'], [['Phạm vi', 'Bộ lọc công nghệ liên quan cả ba nhóm 3,5,7; nguồn từ danh mục 41 trường.'], ['Kết quả', f'{len(strong_techs)} họ công nghệ; {len(support_works)} công trình minh chứng được giữ; {stats["direct_triple_group_works"]} công trình có đủ ba nhãn trực tiếp.'], ['Cách đọc', 'Một họ có thể có bằng chứng ở ba nhóm qua nhiều công trình. Cột ID công trình đủ ba nhãn ghi riêng những công trình trực tiếp đáp ứng cả ba.'], ['Phân biệt G7', 'Ghi rõ năng lượng khối suy luận, năng lượng vận hành, hoặc tối ưu mô hình/hệ thống kỹ thuật. Không suy tổng điện năng robot từ năng lượng chip.'], ['Mức đọc', 'Abstract, metadata và phần chọn lọc theo từng mục; không tuyên bố đã thẩm định toàn văn.']], {1:28, 2:145})
sheet('Cong_nghe_da_loc', techheaders, [techrow(t) for t in strong_techs], techwidths)
sheet('Cong_trinh_minh_chung', headers, [workrow(w) for w in support_works], widths)
sheet('Truc_tiep_du_3_nhom', headers, [workrow(w) for w in works if all(g in w['group_codes'] for g in GROUPS)], widths)
sheet('Chua_du_bang_chung', techheaders, [techrow(t) for t in techs if not t.get('strong_filter')], techwidths)
wb.save(OUT / 'cong_nghe_giao_3_nhom.xlsx')
wb = fullwb

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.header.paragraphs[0].text = 'NGHIÊN CỨU AI • BA NHÓM CHUYÊN ĐỀ 3, 5, 7'
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)
def p(text):
    return doc.add_paragraph(clean(flat(text)))
def hyperlink(paragraph, label, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    co = OxmlElement('w:color')
    co.set(qn('w:val'), '0563C1')
    pr.append(co)
    r.append(pr)
    t = OxmlElement('w:t')
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)

doc.add_heading('Danh mục nghiên cứu AI: nhóm 3, 5 và 7', 0)
p(f'Ngày tra cứu: 30/10/2026. Phạm vi tìm kiếm: {len(scope)} trường tại Việt Nam, Hoa Kỳ và Châu Âu.')
p(f'{len(works)} công trình duy nhất có nguồn nội dung/metadata chuyên biệt được kiểm tra, trong đó {new_count} mục chưa có trong bộ 87 công trình AI trước. Có thêm {len(leads)} đầu mối ở mức thư mục cần đối chiếu tiếp. Một công trình có thể thuộc nhiều nhóm hoặc nhiều trường.')
doc.add_heading('Phạm vi, cách chọn và giới hạn', 1)
for text in rules + limits:
    p(text)
doc.add_heading('Công nghệ đã lọc theo giao ba nhóm', 1)
p(f'Giữ {len(strong_techs)} họ công nghệ, liên kết với {len(support_works)} công trình minh chứng. Trong đó có {stats["direct_triple_group_works"]} công trình được phân loại trực tiếp ở cả ba nhóm; các họ còn lại dựa trên bằng chứng từ nhiều công trình.')
p('G7 gồm cả tối ưu hệ thống kỹ thuật. Mỗi mục phân biệt năng lượng suy luận, năng lượng vận hành robot và giảm chi phí mô hình; không coi các kết quả đó là tương đương.')
for t in strong_techs:
    doc.add_heading(t['technology_id'] + '. ' + t['name_vi'], 2)
    p('Công nghệ cụ thể: ' + '; '.join(t['concrete_technologies']))
    for label, key in [('G3', 'role_g3_vi'), ('G5', 'role_g5_vi'), ('G7', 'role_g7_vi'), ('Cơ chế', 'mechanism_vi'), ('Bằng chứng', 'evidence_note_vi'), ('Mức triển khai / giới hạn', 'maturity_note_vi')]:
        p(label + ': ' + t[key])
    p('ID minh chứng: ' + ', '.join(t['all_work_ids']) + '. ID đủ ba nhãn trực tiếp: ' + (', '.join(t['direct_triple_work_ids']) or 'Bằng chứng qua nhiều công trình.'))
    for url in t['source_urls']:
        par = doc.add_paragraph()
        hyperlink(par, url, url)
doc.add_heading('Kết quả theo nhóm', 1)
for g, title in GROUPS.items():
    doc.add_heading(g + '. ' + title, 2)
    selected = [w for w in works if g in w['group_codes']]
    p(f'{len(selected)} công trình; {stats["group_school_counts"][g]} trường có mục được ghi nhận trong danh mục.')
    for w in selected:
        p(f'{w["work_id"]} | {", ".join(w["schools"])} | {w.get("year") or "Chưa rõ năm"} | {w["title"]}')

doc.add_heading('Nội dung chi tiết theo trường', 1)
printed = set()
for s in scope:
    doc.add_heading(s['name'], 2)
    par = doc.add_paragraph('Website: ')
    hyperlink(par, s['website'], s['website'])
    relevant = [w for w in works if s['short_name'] in w['schools']]
    if not relevant:
        p('Đã có nhật ký truy vấn cho ba nhóm; chưa tìm được công trình đủ nguồn trong phạm vi lượt tìm. Điều này không chứng minh trường không nghiên cứu các nhóm này.')
    for w in relevant:
        if w['work_id'] in printed:
            p(f'{w["work_id"]}: {w["title"]}. Công trình hợp tác đã trình bày ở trường trước; xem ID trong Excel để tra đủ bằng chứng.')
            continue
        printed.add(w['work_id'])
        doc.add_heading(w['work_id'] + '. ' + w['title'], 3)
        p(f'Trường: {", ".join(w["schools"])} | Năm: {w.get("year") or "Chưa xác định"} | Loại: {w.get("type")} | Trạng thái: {w["status"]}')
        p('Nhóm: ' + '; '.join(w['group_names_vi']) + '. Chủ đề cụ thể: ' + '; '.join(w['subtopics_vi']))
        for label, key in [('Bài toán', 'problem_vi'), ('Phương pháp', 'method_vi'), ('Đóng góp / mục tiêu theo nguồn', 'contribution_vi'), ('Ứng dụng', 'application_vi')]:
            if w.get(key):
                p(label + ': ' + flat(w[key]))
        p('Bằng chứng quan hệ trường: ' + w['institution_evidence_vi'])
        p('Mức đọc: ' + w.get('reading_depth', '') + '. Kiểm chứng: ' + ' '.join(w['verification_notes_vi']))
        if w.get('doi'):
            p('DOI: ' + w['doi'])
        for url in w['source_urls']:
            par = doc.add_paragraph()
            hyperlink(par, url, url)
        for label, key in [('Mã nguồn', 'code_url'), ('Dữ liệu', 'data_url')]:
            if w.get(key):
                par = doc.add_paragraph(label + ': ')
                hyperlink(par, str(w[key]), str(w[key]))

doc.add_heading('Phụ lục: đầu mối mới ở mức thư mục', 1)
p('Các mục dưới đây chưa có đủ đối chiếu nội dung để cộng vào danh mục chính. Năm và tên tác giả có thể cần sửa sau khi đối chiếu nhà xuất bản.')
for r in leads:
    p(f'{r["lead_id"]} | {flat(r.get("school", r.get("schools")))} | {r.get("year")} | {r["title"]} | Nhóm dự kiến: {", ".join(r["group_codes"])}')
    for url in unique(r.get('lead_sources', [])):
        par = doc.add_paragraph()
        hyperlink(par, url, url)
doc.add_heading('Nhật ký tìm kiếm', 1)
p(f'{len(audits)} bản ghi trường–nhóm được lưu đầy đủ trong Excel và JSON. Mỗi bản ghi giữ truy vấn, nguồn và kết luận có giới hạn, bao gồm các trường chưa tìm được công trình phù hợp.')
doc.save(OUT / 'bao_cao_chi_tiet_AI_3_5_7.docx')

(OUT / 'HUONG_DAN.txt').write_text('\n'.join([
    'NGHIEN CUU AI NHOM 3, 5, 7 - 30/10/2026',
    f'{len(works)} cong trinh chinh; {len(leads)} dau moi thu muc; pham vi41truong.',
    'Excel: loc theo truong, nhom, chu de, nam, trang thai; cac sheet nhom co ID trung vi cong trinh co the nhieu nhan.',
    'Word: danh muc theo nhom va mo ta tung cong trinh. JSON/CSV: du lieu de truy van.',
    'Xem Huong_dan, Dau_moi_can_doi_chieu va Kiem_chung_truong de danh gia muc bang chung va pham vi.',
    *limits,
]), encoding='utf-8')
zip_path = ROOT / 'bao_cao_AI_nhom_3_5_7.zip'
with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
    for path in sorted(OUT.iterdir()):
        if path.is_file():
            z.write(path, path.name)
print(json.dumps(stats, ensure_ascii=False, indent=2))
print(str(zip_path))
