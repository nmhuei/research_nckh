import sys,json,re,csv,zipfile,collections
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'.deps'))
from openpyxl import load_workbook
from docx import Document
from config import ROOT,TOPICS,SEEDS,ALIASES

out=ROOT/'deliverables'
data=json.loads((out/'research_classified.json').read_text());works=data['research_items'];schools=data['schools']
stats=json.loads((out/'statistics.json').read_text());bib=json.loads((out/'bibliography_technology_only.json').read_text())['records']
key=lambda s:re.sub(r'\s+',' ',s).strip().casefold()
all_works=[w for n in ['vn_core_research','vn_research','us_research','eu_research'] for s in json.loads((ROOT/'parts'/f'{n}.json').read_text())['schools'] for w in s['items']]
all_bib=json.loads((ROOT/'bibliography_classified.json').read_text())['records']
def read_decisions(names):
 keep=set();removed=set();total=[]
 for name in names:
  d=json.loads((ROOT/'parts'/f'{name}.json').read_text())
  k=[key(v if isinstance(v,str) else v['title']) for v in d['keep']]
  x=[key(v['title']) for field in ['exclude','uncertain'] for v in d[field]]
  assert not (set(k)&set(x)),name
  keep.update(k);removed.update(x);total.extend(k+x)
 return keep,removed,total
wk,wx,wt=read_decisions(['vn_relevance_review','us_relevance_review','eu_relevance_review'])
bk,bx,bt=read_decisions(['hust_bib_relevance','vku_bib_relevance','ptit_bib_relevance'])
assert {key(w['title']) for w in all_works}==set(wt)
assert {key(w['title']) for w in all_bib}==set(bt)
assert {key(w['title']) for w in works}==wk
assert {key(w['title']) for w in bib}==bk
assert not ({key(w['title']) for w in works}&wx)
assert not ({key(w['title']) for w in bib}&bx)
assert len(works)==stats['selected_work_records']
assert len(bib)==stats['bibliography_kept']
assert sum(stats['primary_topic_counts'].values())==len(works)
assert len({w['record_id'] for w in works})==len(works)
for w in works+bib:
 assert w['topic_codes'],w['title']
 assert set(w['topic_codes'])<=set(TOPICS)-{'OTHER','BASIC_SCIENCE'},w['title']
 assert w['primary_topic_code'] in w['topic_codes'],w['title']
 if w.get('doi'):assert re.match(r'^10\.\d{4,9}/\S+',w['doi']),w['doi']
for w in works:
 assert w['url'].startswith(('https://','http://'))
 assert w.get('verification_note') and w.get('summary_vi'),w['title']
 assert w['year'] is None or 1900<=w['year']<=2026,w['title']
assert next(w for w in works if w['title'].startswith('PRADA-QA'))['year']==2026
with (out/'research_classified.csv').open(encoding='utf-8-sig') as f:assert len(list(csv.reader(f)))==len(works)+1
wb=load_workbook(out/'bao_cao_truong_cong_nghe.xlsx',read_only=True,data_only=False)
assert wb['Cong_trinh'].max_row==len(works)+1
assert wb['Truong_cong_nghe'].max_row==len(schools)+1
assert wb['Thu_muc_da_loc'].max_row==len(bib)+1
assert 'Danh_ba_the_gioi' not in wb.sheetnames
assert sum(c.data_type=='f' for ws in wb for row in ws for c in row)==0
wb.close()
doc=Document(out/'bao_cao_nghien_cuu_chi_tiet.docx')
headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 3']
assert len(headings)==len(works)
assert all(any(w['title'] in h for h in headings) for w in works)
with zipfile.ZipFile(ROOT/'bao_cao_truong_cong_nghe.zip') as z:
 assert z.testzip() is None
 assert set(z.namelist())==set(p.name for p in out.iterdir())
 assert 'university_websites.csv' not in z.namelist()
 for n in z.namelist():
  if n.endswith('.json'):json.loads(z.read(n))
result={'status':'PASS','schools':len(schools),'technology_work_records':len(works),'technology_bibliography_records':len(bib),
 'removed_detailed':len(all_works)-len(works),'removed_bibliography':len(all_bib)-len(bib),
 'primary_topics':len(stats['primary_topic_counts']),'works_with_doi':stats['works_with_doi'],'xlsx_docx_zip_readable':True,'excluded_titles_absent':True,
 'scope_complete_worldwide':False,'full_text_review_complete':False}
(ROOT/'validation_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
