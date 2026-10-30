import collections
import csv
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
OUT = BASE / 'deliverables'
sys.path.insert(0, str(ROOT / '.deps'))
from openpyxl import load_workbook
from docx import Document

def normtitle(t):
    t = unicodedata.normalize('NFKD', t).casefold()
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in t if not unicodedata.combining(c)))

d = json.loads((OUT / 'catalogue_AI_3_5_7.json').read_text())
works, schools = d['works'], d['schools']
stats = d['metadata']['statistics']
groups = set(d['metadata']['groups'])
schoolnames = {s['short_name'] for s in schools}
ids = {w['work_id'] for w in works}
byid = {w['work_id']: w for w in works}
assert len(works) == stats['unique_works'] == len(ids)
assert len(schools) == 41
assert len({normtitle(w['title']) for w in works}) == len(works)
dois = [w['doi'] for w in works if w.get('doi')]
assert len(dois) == len(set(dois)) == stats['works_with_doi']
for w in works:
    for field in ['title', 'schools', 'group_codes', 'type', 'status', 'problem_vi', 'method_vi', 'contribution_vi', 'application_vi', 'url', 'source_urls', 'institution_evidence_vi', 'reading_depth', 'verification_notes_vi']:
        assert w.get(field), (w['work_id'], field)
    assert set(w['schools']) <= schoolnames
    assert set(w['group_codes']) <= groups
    assert isinstance(w.get('year'), int) and 2000 <= w['year'] <= 2026
    assert all(isinstance(u, str) and u.startswith(('https://', 'http://')) for u in w['source_urls'])
    assert not w['full_text_read']
    assert len(w['evidence_records']) >= 1

assert len(d['search_audit']) == stats['school_group_search_audit_records'] == 123
for g in groups:
    audits = [a for a in d['search_audit'] if a['group_code'] == g]
    assert {a['school'] for a in audits} == schoolnames
    assert len(audits) == 41
    assert all(a.get('queries') and a.get('result_note_vi') for a in audits)
    assert stats['group_counts'][g] == sum(g in w['group_codes'] for w in works)
assert len(d['previous_bibliography_screening']) == 287
assert len(d['bibliographic_leads']) == stats['bibliographic_leads_not_in_main_catalogue']
assert stats['school_work_associations'] == sum(len(w['schools']) for w in works)
assert stats['direct_triple_group_works'] == sum(set(w['group_codes']) == groups for w in works) == 3

filtered = json.loads((OUT / 'cong_nghe_giao_3_nhom.json').read_text())
techs = filtered['technologies']
assert len(techs) == stats['technology_families_retained_across_three_groups'] == 5
all_support = set()
for t in techs:
    assert t['strong_filter']
    assert not t['unmatched_titles'], (t['technology_id'], t['unmatched_titles'])
    for g in groups:
        assert t['work_ids_by_group'][g]
        for i in t['work_ids_by_group'][g]:
            assert i in ids and g in byid[i]['group_codes'], (t['technology_id'], g, i)
    assert t['direct_triple_work_ids'] or not t['single_paper_all_three_confirmed']
    for i in t['direct_triple_work_ids']:
        assert set(byid[i]['group_codes']) == groups
    all_support.update(t['all_work_ids'])
assert {w['work_id'] for w in filtered['supporting_works']} == all_support
assert len(all_support) == stats['technology_supporting_works']
assert len(filtered['not_retained']) == 1
assert not filtered['not_retained'][0]['work_ids_by_group']['G5']

titlemap = {normtitle(w['title']): w for w in works}
assert titlemap[normtitle('Toward Sustainable and Cost-Efficient HPC Systems: A Deep Reinforcement Learning Job Scheduling Approach')]['group_codes'] == ['G7']
assert titlemap[normtitle('LLMs Can Reason Faster Only If We Let Them')]['group_codes'] == ['G3']
assert titlemap[normtitle('Early Fault Detection in Pump Systems Using Unsupervised Anomaly Detection Techniques')]['year'] == 2026
assert titlemap[normtitle('DPaI: Differentiable Pruning at Initialization with Node-Path Balance Principle')]['schools'] == ['HUST']

with (OUT / 'catalogue_AI_3_5_7.csv').open(encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))
assert len(rows) == len(works) + 1
wb = load_workbook(OUT / 'bao_cao_AI_nhom_3_5_7.xlsx', data_only=False)
assert wb['Tat_ca_cong_trinh'].max_row == len(works) + 1
assert wb['Kiem_chung_truong'].max_row == 124
assert wb['Theo_truong'].max_row == 42
for group, name in [('G3', 'G3_AI_nhe_va_thiet_bi'), ('G5', 'G5_Robot_tu_hanh'), ('G7', 'G7_Nang_luong_ky_thuat')]:
    assert wb[name].max_row == stats['group_counts'][group] + 1
small = load_workbook(OUT / 'cong_nghe_giao_3_nhom.xlsx', data_only=False)
assert small['Cong_nghe_da_loc'].max_row == 6
assert small['Cong_trinh_minh_chung'].max_row == len(all_support) + 1
assert small['Truc_tiep_du_3_nhom'].max_row == 4
for workbook in [wb, small]:
    for ws in workbook:
        for row in ws:
            for c in row:
                assert c.data_type != 'f', (ws.title, c.coordinate)
doc = Document(OUT / 'bao_cao_chi_tiet_AI_3_5_7.docx')
text = '\n'.join(p.text for p in doc.paragraphs)
for w in works:
    assert w['work_id'] + '. ' + w['title'] in text
for t in techs:
    assert t['technology_id'] + '. ' + t['name_vi'] in text
with zipfile.ZipFile(OUT / 'bao_cao_chi_tiet_AI_3_5_7.docx') as z:
    ElementTree.fromstring(z.read('word/document.xml'))
with zipfile.ZipFile(ROOT / 'bao_cao_AI_nhom_3_5_7.zip') as z:
    assert z.testzip() is None
    assert set(z.namelist()) == {p.name for p in OUT.iterdir() if p.is_file()}

result = {'result': 'PASS', 'date': '2026-10-30', 'unique_works': len(works), 'technology_families_retained': len(techs), 'supporting_works': len(all_support), 'direct_triple_group_works': 3, 'scope_schools': 41, 'school_group_search_records': 123, 'group_counts': stats['group_counts'], 'bibliographic_leads': len(d['bibliographic_leads']), 'file_formats_checked': ['JSON', 'CSV', 'XLSX', 'DOCX', 'ZIP']}
(BASE / 'validation_result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False, indent=2))
